"""Turning a stored conversion into a downloadable file.

Strictly a read and a reformat. This module never calls a model, never re-extracts, and
never invokes ``ConversionService``: it takes a conversion that was already produced and
persisted, and writes it out in the requested format. A download is therefore free,
instant, and — importantly — identical every time, because it is derived from stored
bytes rather than from a fresh generation that might come back different.

Four formats, and each exists for a different reader:

- **Markdown** is returned from the stored text verbatim. It was rendered once, when the
  conversion ran, and re-rendering it here would risk a download that disagrees with what
  the user was shown.
- **DOCX** is the same document as a Word file, with real headings and real tables.
- **XLSX** is the data form: one worksheet per section, so a reader can sort and filter.
- **JSON** is the whole result including the correction audit, for anything downstream.

No new dependency was added for any of them: ``python-docx`` and ``openpyxl`` were already
in the project for *parsing* uploads, and both write as well as read.

Two safety rules carry over from the existing export path. Values written into a
spreadsheet are escaped against formula injection, since a cell beginning ``=``, ``+``,
``-`` or ``@`` is executed by Excel and the values here came from a user's document. And
filenames go through the same ``sanitize_filename`` every other download uses.
"""

from __future__ import annotations

import io
import json
from collections.abc import Callable
from typing import Literal

from docx import Document as DocxDocument
from docx.document import Document as DocxDocumentType
from openpyxl import Workbook
from openpyxl.styles import Alignment

from app.api.schemas.conversion import (
    ConversionResponse,
    ConvertedFieldResponse,
    ConvertedRecordResponse,
    ConvertedSectionResponse,
)
from app.core.security import sanitize_filename
from app.models.domain.conversion import SectionKind, ValueOrigin

# Reused rather than reimplemented: this is the project's single definition of how a
# spreadsheet cell is neutralised, and a second copy would be one edit away from
# disagreeing with the first. It is underscore-prefixed only because the structure
# exporter had no other caller when it was written.
from app.services.export_service import ExportedFile, _escape_cell

ConversionExportFormat = Literal["md", "docx", "xlsx", "json"]

_MEDIA_TYPES: dict[ConversionExportFormat, str] = {
    "md": "text/markdown; charset=utf-8",
    "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    "json": "application/json",
}

#: Excel rejects these in a worksheet name, and caps the name at 31 characters.
_INVALID_SHEET_CHARS = set(r"[]:*?/\'")
_MAX_SHEET_NAME = 31

#: Applied to values whose line structure matters. A typeface, not a language: nothing
#: here knows or cares which language a block of code is written in, only that its lines
#: and indentation should survive.
_MONOSPACE_FONT = "Consolas"


def export_conversion(
    conversion: ConversionResponse, export_format: ConversionExportFormat
) -> ExportedFile:
    """Render a stored conversion as a downloadable file.

    Args:
        conversion: The persisted result, exactly as it was produced.
        export_format: Which format to write.

    Returns:
        The file's bytes, its media type, and a sanitised filename.
    """
    content = _FORMATTERS[export_format](conversion)
    filename = sanitize_filename(f"{conversion.title or 'converted-document'}.{export_format}")
    return ExportedFile(content=content, media_type=_MEDIA_TYPES[export_format], filename=filename)


def _to_markdown(conversion: ConversionResponse) -> bytes:
    """The stored rendering, encoded.

    Deliberately not re-rendered. The Markdown was produced when the conversion ran;
    generating it again here would mean a download could differ from the preview the user
    approved, for no benefit.
    """
    return conversion.markdown.encode("utf-8")


def _to_json(conversion: ConversionResponse) -> bytes:
    """The whole result, including the correction audit.

    ``ensure_ascii=False`` so a name written in a non-Latin script survives as itself
    rather than as escape sequences. Never escaped for formula injection — a JSON parser
    does not evaluate a leading ``=``, and escaping would corrupt the value.
    """
    payload = conversion.model_dump(mode="json")
    return json.dumps(payload, indent=2, ensure_ascii=False).encode("utf-8")


def _to_docx(conversion: ConversionResponse) -> bytes:
    """The document as a Word file: real headings, real tables."""
    document: DocxDocumentType = DocxDocument()
    document.add_heading(conversion.title or "Converted Document", level=0)

    for section in conversion.sections:
        document.add_heading(section.heading, level=1)

        if section.kind is SectionKind.NOTE:
            if section.note.strip():
                document.add_paragraph(section.note.strip())
            continue

        if not section.records:
            document.add_paragraph("No records in this section.")
            continue

        if section.note.strip():
            document.add_paragraph(section.note.strip())

        if section.kind is SectionKind.DETAIL:
            _write_docx_detail(document, section)
        else:
            _write_docx_table(document, section)

    _write_docx_audit(document, conversion)

    if conversion.warnings:
        document.add_heading("Warnings", level=1)
        for warning in conversion.warnings:
            document.add_paragraph(warning, style="List Bullet")

    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _write_docx_table(document: DocxDocumentType, section: ConvertedSectionResponse) -> None:
    labels = _column_labels(section)
    table = document.add_table(rows=1, cols=len(labels))
    table.style = "Light Grid Accent 1"

    for cell, label in zip(table.rows[0].cells, labels, strict=True):
        cell.text = label

    for record in section.records:
        values = _record_values(record, section.columns)
        row = table.add_row()
        for cell, value in zip(row.cells, values, strict=True):
            cell.text = value


def _write_docx_detail(document: DocxDocumentType, section: ConvertedSectionResponse) -> None:
    for record in section.records:
        for name in section.columns:
            item = _field(record, name)
            if item is None:
                continue
            text = _display_text(item)
            source = _source_label(item)
            suffix = f"  (source: {source})" if source else ""

            if _is_multiline(text):
                # Label and value are separated so the value can be a block. python-docx
                # already turns "\n" into a line break and marks leading spaces with
                # xml:space="preserve", so the text arrives in Word with its lines and
                # indentation intact; what it does not do is make it monospaced, which
                # is what the run below adds. The string itself is passed through
                # untouched.
                document.add_paragraph(f"{item.label}:{suffix}", style="List Bullet")
                run = document.add_paragraph().add_run(text)
                run.font.name = _MONOSPACE_FONT
                continue

            document.add_paragraph(f"{item.label}: {text}{suffix}", style="List Bullet")
        document.add_paragraph()


def _write_docx_audit(document: DocxDocumentType, conversion: ConversionResponse) -> None:
    """The corrections audit, so the Word file explains its own changes.

    A downloaded report that has been silently tidied is exactly what this feature exists
    to avoid, so the audit travels with every format that can carry it.
    """
    if not conversion.corrections and not conversion.suggestions:
        return

    document.add_heading("Changes & Corrections", level=1)

    for item in conversion.corrections:
        document.add_paragraph(
            f'"{item.original}" → "{item.corrected}" — {item.reason} '
            f"(confidence: {item.confidence.value}, applied)",
            style="List Number",
        )

    if conversion.suggestions:
        document.add_heading("Possible corrections (not applied)", level=2)
        for item in conversion.suggestions:
            document.add_paragraph(
                f'"{item.original}" → "{item.corrected}" — {item.reason} '
                f"(confidence: {item.confidence.value}, NOT applied; "
                "the original value is shown)",
                style="List Number",
            )


def _to_xlsx(conversion: ConversionResponse) -> bytes:
    """The data form: one worksheet per section, plus the audit.

    Every cell goes through the same formula-injection escaping the structure export
    uses — the values came from a user's document and are untrusted exactly like any
    other document content.
    """
    workbook = Workbook()
    workbook.remove(workbook.active)

    used: set[str] = set()
    for section in conversion.sections:
        sheet = workbook.create_sheet(_sheet_name(section.heading, used))

        if section.kind is SectionKind.NOTE:
            sheet.append(["Note"])
            sheet.append([_escape_cell(section.note.strip())])
            continue

        sheet.append([_escape_cell(label) for label in _column_labels(section)])
        for record in section.records:
            values = _record_values(record, section.columns)
            sheet.append([_escape_cell(value) for value in values])
            # Excel stores the newlines regardless but renders only the first line
            # unless the cell wraps, so a multi-line value would look truncated.
            for column, value in enumerate(values, start=1):
                if _is_multiline(value):
                    sheet.cell(row=sheet.max_row, column=column).alignment = Alignment(
                        wrap_text=True, vertical="top"
                    )

    audit = workbook.create_sheet(_sheet_name("Changes & Corrections", used))
    audit.append(["original", "corrected", "reason", "confidence", "applied"])
    for item in [*conversion.corrections, *conversion.suggestions]:
        audit.append(
            [
                _escape_cell(item.original),
                _escape_cell(item.corrected),
                _escape_cell(item.reason),
                item.confidence.value,
                item.applied,
            ]
        )

    if not workbook.sheetnames:  # pragma: no cover - the audit sheet is always added
        workbook.create_sheet("Empty")

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _sheet_name(heading: str, used: set[str]) -> str:
    """A worksheet name Excel will accept, and that has not been used already.

    Excel forbids several characters, caps names at 31 chars, and rejects duplicates —
    all three of which a section heading can trip over, since headings are free text and
    a grouped plan can produce two sections whose names collide once truncated.
    """
    cleaned = "".join(" " if ch in _INVALID_SHEET_CHARS else ch for ch in heading).strip()
    cleaned = " ".join(cleaned.split()) or "Section"
    cleaned = cleaned[:_MAX_SHEET_NAME]

    candidate = cleaned
    suffix = 2
    while candidate.casefold() in used:
        tail = f" ({suffix})"
        candidate = cleaned[: _MAX_SHEET_NAME - len(tail)] + tail
        suffix += 1

    used.add(candidate.casefold())
    return candidate


def _column_labels(section: ConvertedSectionResponse) -> list[str]:
    """Column headers, taken from the first record that carries each field."""
    labels: list[str] = []
    for name in section.columns:
        label = name
        for record in section.records:
            item = _field(record, name)
            if item is not None:
                label = item.label
                break
        labels.append(label)
    return labels


def _record_values(record: ConvertedRecordResponse, columns: list[str]) -> list[str]:
    values: list[str] = []
    for name in columns:
        item = _field(record, name)
        values.append("" if item is None else _display_text(item))
    return values


def _field(record: ConvertedRecordResponse, name: str) -> ConvertedFieldResponse | None:
    for item in record.fields:
        if item.name == name:
            return item
    return None


def _display_text(item: ConvertedFieldResponse) -> str:
    """The value as a reader should see it. An absent value is blank, never invented."""
    if item.value.origin is ValueOrigin.ABSENT:
        return ""
    return item.value.text


def _is_multiline(text: str) -> bool:
    """Whether the source gave this value more than one line.

    Mirrors ``conversion_render._is_multiline`` deliberately: both answer the same
    question about the same text, and an export that disagreed with the preview about
    which values are blocks would be its own bug.
    """
    return "\n" in text or "\r" in text


def _source_label(item: ConvertedFieldResponse) -> str:
    if not item.value.is_verified or not item.value.source_label:
        return ""
    return item.value.source_label


_FORMATTERS: dict[ConversionExportFormat, Callable[[ConversionResponse], bytes]] = {
    "md": _to_markdown,
    "docx": _to_docx,
    "xlsx": _to_xlsx,
    "json": _to_json,
}
