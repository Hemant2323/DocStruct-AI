"""Tests for the conversion export formatters.

Pure logic, no network and no database — each formatter is a function from a stored
conversion to bytes, so it can be tested by handing it one.

The properties that matter: every record reaches every format, a value that came from a
user's document cannot execute in a spreadsheet, non-Latin text survives, and the same
input always produces the same file.
"""

from __future__ import annotations

import io
import json
from datetime import UTC, datetime
from uuid import uuid4

import pytest
from docx import Document as DocxDocument
from openpyxl import load_workbook

from app.api.schemas.conversion import (
    ConversionCorrectionResponse,
    ConversionPlanResponse,
    ConversionResponse,
    ConvertedFieldResponse,
    ConvertedRecordResponse,
    ConvertedSectionResponse,
    ConvertedValueResponse,
)
from app.models.domain.conversion import Confidence, CorrectionKind, SectionKind, ValueOrigin
from app.services.conversion_export import export_conversion

DOCUMENT_ID = uuid4()
CONVERSION_ID = uuid4()


def _value(
    text: str,
    *,
    origin: ValueOrigin = ValueOrigin.SOURCE,
    source_label: str | None = "page 1",
    is_verified: bool = True,
) -> ConvertedValueResponse:
    return ConvertedValueResponse(
        text=text,
        origin=origin,
        source_label=source_label if is_verified else None,
        quoted_text=text if is_verified else None,
        is_verified=is_verified,
    )


def _field(name: str, label: str, value: ConvertedValueResponse) -> ConvertedFieldResponse:
    return ConvertedFieldResponse(name=name, label=label, value=value)


def _record(index: int, holder: str, bank: str) -> ConvertedRecordResponse:
    return ConvertedRecordResponse(
        source_record_index=index,
        fields=[
            _field("account_holder", "Account Holder", _value(holder)),
            _field("bank_name", "Bank", _value(bank)),
        ],
    )


def _section(
    *records: ConvertedRecordResponse, heading: str = "Customers"
) -> ConvertedSectionResponse:
    return ConvertedSectionResponse(
        heading=heading,
        kind=SectionKind.TABLE,
        columns=["account_holder", "bank_name"],
        records=list(records),
    )


def _correction(applied: bool = True) -> ConversionCorrectionResponse:
    return ConversionCorrectionResponse(
        id="c1",
        kind=CorrectionKind.VOCABULARY,
        record_index=0,
        field="account_type",
        original="sav",
        corrected="Savings",
        reason="standardized account type",
        confidence=Confidence.HIGH if applied else Confidence.MEDIUM,
        applied=applied,
    )


def _conversion(**kwargs: object) -> ConversionResponse:
    defaults: dict[str, object] = {
        "id": CONVERSION_ID,
        "document_id": DOCUMENT_ID,
        "title": "Customer Master Report",
        "instruction": "Create a clean customer master report.",
        "markdown": "# Customer Master Report\n\n## Customers\n",
        "sections": [
            _section(
                _record(0, "Rajesh Kumar Sharma", "State Bank of India"),
                _record(1, "Priya Sharma", "HDFC Bank"),
                _record(2, "Rahul Verma", "ICICI Bank"),
            )
        ],
        "plan": ConversionPlanResponse(title="Customer Master Report"),
        "corrections": [_correction()],
        "suggestions": [],
        "rejected": [],
        "schema_used": {"account_holder": "string", "bank_name": "string"},
        "vocabularies": {},
        "source_record_count": 3,
        "warnings": [],
        "model": "fake-model",
        "input_tokens": 100,
        "output_tokens": 50,
        "records_from_cache": False,
        "created_at": datetime(2026, 8, 26, 12, 0, tzinfo=UTC),
    }
    defaults.update(kwargs)
    return ConversionResponse.model_validate(defaults)


class TestMarkdownExport:
    def test_the_stored_rendering_is_returned_verbatim(self) -> None:
        """Never re-rendered: a download must not disagree with the preview."""
        conversion = _conversion(markdown="# Title\n\n| A | B |\n")

        exported = export_conversion(conversion, "md")

        assert exported.content.decode("utf-8") == "# Title\n\n| A | B |\n"
        assert exported.media_type.startswith("text/markdown")

    def test_unicode_survives(self) -> None:
        conversion = _conversion(markdown="# शीर्षक\n\nराजेश — ₹1,45,320\n")

        exported = export_conversion(conversion, "md")

        assert "शीर्षक" in exported.content.decode("utf-8")
        assert "₹1,45,320" in exported.content.decode("utf-8")


class TestJsonExport:
    def test_every_record_is_present(self) -> None:
        payload = json.loads(export_conversion(_conversion(), "json").content)

        holders = [
            item["value"]["text"]
            for section in payload["sections"]
            for record in section["records"]
            for item in record["fields"]
            if item["name"] == "account_holder"
        ]
        assert holders == ["Rajesh Kumar Sharma", "Priya Sharma", "Rahul Verma"]

    def test_the_correction_audit_is_included(self) -> None:
        conversion = _conversion(corrections=[_correction()], suggestions=[_correction(False)])

        payload = json.loads(export_conversion(conversion, "json").content)

        assert payload["corrections"][0]["corrected"] == "Savings"
        assert payload["suggestions"][0]["applied"] is False

    def test_unicode_is_not_escaped(self) -> None:
        conversion = _conversion(sections=[_section(_record(0, "राजेश शर्मा", "स्टेट बैंक"))])

        raw = export_conversion(conversion, "json").content.decode("utf-8")

        assert "राजेश शर्मा" in raw
        assert "\\u0930" not in raw

    def test_values_are_not_formula_escaped(self) -> None:
        """A JSON parser does not evaluate a leading '='; escaping would corrupt it."""
        conversion = _conversion(sections=[_section(_record(0, "=SUM(A1:A2)", "Bank"))])

        payload = json.loads(export_conversion(conversion, "json").content)

        holder = payload["sections"][0]["records"][0]["fields"][0]["value"]["text"]
        assert holder == "=SUM(A1:A2)"

    def test_no_internal_identifiers_are_exposed(self) -> None:
        raw = export_conversion(_conversion(), "json").content.decode("utf-8")

        assert "chunk_id" not in raw
        assert "block_ids" not in raw


class TestDocxExport:
    def _document(self, conversion: ConversionResponse):  # type: ignore[no-untyped-def]
        return DocxDocument(io.BytesIO(export_conversion(conversion, "docx").content))

    def test_the_title_and_headings_are_real_word_styles(self) -> None:
        """Not bold paragraphs: the title carries Word's 'Title' style and sections carry
        'Heading N', so the file has a real outline a reader can navigate."""
        document = self._document(_conversion())

        styled = {p.text: p.style.name for p in document.paragraphs if p.text}
        assert styled["Customer Master Report"] == "Title"
        assert styled["Customers"].startswith("Heading")

    def test_every_record_becomes_a_table_row(self) -> None:
        document = self._document(_conversion())

        [table] = [t for t in document.tables if t.rows[0].cells[0].text == "Account Holder"]
        assert len(table.rows) == 4  # header + 3 records
        assert table.rows[1].cells[0].text == "Rajesh Kumar Sharma"
        assert table.rows[3].cells[1].text == "ICICI Bank"

    def test_column_headers_use_labels_not_field_keys(self) -> None:
        document = self._document(_conversion())

        header = [cell.text for cell in document.tables[0].rows[0].cells]
        assert header == ["Account Holder", "Bank"]

    def test_the_correction_audit_is_written(self) -> None:
        conversion = _conversion(corrections=[_correction()], suggestions=[_correction(False)])

        text = "\n".join(p.text for p in self._document(conversion).paragraphs)

        assert "Changes & Corrections" in text
        assert '"sav" → "Savings"' in text
        assert "NOT applied" in text

    def test_unicode_survives(self) -> None:
        conversion = _conversion(
            title="ग्राहक रिपोर्ट",
            sections=[_section(_record(0, "राजेश शर्मा", "स्टेट बैंक"))],
        )
        document = self._document(conversion)

        text = "\n".join(p.text for p in document.paragraphs)
        assert "ग्राहक रिपोर्ट" in text
        assert document.tables[0].rows[1].cells[0].text == "राजेश शर्मा"

    def test_an_absent_value_is_blank_not_invented(self) -> None:
        record = ConvertedRecordResponse(
            source_record_index=0,
            fields=[
                _field("account_holder", "Account Holder", _value("Priya Sharma")),
                _field(
                    "bank_name",
                    "Bank",
                    _value("", origin=ValueOrigin.ABSENT, is_verified=False),
                ),
            ],
        )
        document = self._document(_conversion(sections=[_section(record)]))

        assert document.tables[0].rows[1].cells[1].text == ""

    def test_an_empty_section_says_so(self) -> None:
        empty = ConvertedSectionResponse(
            heading="Pending KYC", kind=SectionKind.TABLE, columns=["account_holder"]
        )
        text = "\n".join(p.text for p in self._document(_conversion(sections=[empty])).paragraphs)

        assert "No records in this section." in text

    def test_a_note_section_is_written_as_prose(self) -> None:
        note = ConvertedSectionResponse(
            heading="Notes", kind=SectionKind.NOTE, note="Prepared for review."
        )
        text = "\n".join(p.text for p in self._document(_conversion(sections=[note])).paragraphs)

        assert "Prepared for review." in text

    def test_warnings_are_written(self) -> None:
        conversion = _conversion(warnings=["2 corrections were rejected as unsafe."])

        text = "\n".join(p.text for p in self._document(conversion).paragraphs)

        assert "2 corrections were rejected as unsafe." in text


class TestXlsxExport:
    def _workbook(self, conversion: ConversionResponse):  # type: ignore[no-untyped-def]
        return load_workbook(io.BytesIO(export_conversion(conversion, "xlsx").content))

    def test_one_worksheet_per_section_plus_the_audit(self) -> None:
        conversion = _conversion(
            sections=[
                _section(_record(0, "A", "B"), heading="Savings"),
                _section(_record(1, "C", "D"), heading="Current"),
            ]
        )

        workbook = self._workbook(conversion)

        assert workbook.sheetnames == ["Savings", "Current", "Changes & Corrections"]

    def test_every_record_becomes_a_row(self) -> None:
        workbook = self._workbook(_conversion())
        rows = list(workbook["Customers"].iter_rows(values_only=True))

        assert rows[0] == ("Account Holder", "Bank")
        assert len(rows) == 4
        assert rows[3] == ("Rahul Verma", "ICICI Bank")

    @pytest.mark.parametrize("trigger", ["=", "+", "-", "@"])
    def test_formula_injection_is_escaped(self, trigger: str) -> None:
        """Values came from a user's document; a cell starting with these executes."""
        hostile = f"{trigger}cmd|' /C calc'!A0"
        conversion = _conversion(sections=[_section(_record(0, hostile, "Bank"))])

        rows = list(self._workbook(conversion)["Customers"].iter_rows(values_only=True))

        assert rows[1][0] == f"'{hostile}"

    def test_a_normal_value_is_not_escaped(self) -> None:
        rows = list(self._workbook(_conversion())["Customers"].iter_rows(values_only=True))

        assert rows[1][0] == "Rajesh Kumar Sharma"

    def test_the_audit_sheet_carries_applied_and_suggested(self) -> None:
        conversion = _conversion(corrections=[_correction()], suggestions=[_correction(False)])

        rows = list(self._workbook(conversion)["Changes & Corrections"].iter_rows(values_only=True))

        assert rows[0] == ("original", "corrected", "reason", "confidence", "applied")
        assert rows[1][4] is True
        assert rows[2][4] is False

    def test_unicode_survives(self) -> None:
        conversion = _conversion(sections=[_section(_record(0, "राजेश शर्मा", "स्टेट बैंक"))])

        rows = list(self._workbook(conversion)["Customers"].iter_rows(values_only=True))

        assert rows[1] == ("राजेश शर्मा", "स्टेट बैंक")

    def test_an_invalid_sheet_name_is_made_acceptable(self) -> None:
        """Excel forbids these characters and caps names at 31 chars, but a heading is
        free text produced from a grouped plan."""
        conversion = _conversion(
            sections=[_section(_record(0, "A", "B"), heading="Customers/Savings: [2026]*?")]
        )

        names = self._workbook(conversion).sheetnames

        assert all(ch not in names[0] for ch in r"[]:*?/\'")
        assert len(names[0]) <= 31

    def test_colliding_sheet_names_are_disambiguated(self) -> None:
        long_heading = "Customers grouped by account type " + "x" * 20
        conversion = _conversion(
            sections=[
                _section(_record(0, "A", "B"), heading=long_heading),
                _section(_record(1, "C", "D"), heading=long_heading),
            ]
        )

        names = self._workbook(conversion).sheetnames

        assert len(set(names)) == len(names)

    def test_an_absent_value_is_blank(self) -> None:
        record = ConvertedRecordResponse(
            source_record_index=0,
            fields=[
                _field("account_holder", "Account Holder", _value("Priya Sharma")),
                _field(
                    "bank_name", "Bank", _value("", origin=ValueOrigin.ABSENT, is_verified=False)
                ),
            ],
        )
        rows = list(
            self._workbook(_conversion(sections=[_section(record)]))["Customers"].iter_rows(
                values_only=True
            )
        )

        assert rows[1][1] in (None, "")


class TestFilenames:
    @pytest.mark.parametrize(
        ("export_format", "extension"),
        [("md", ".md"), ("docx", ".docx"), ("xlsx", ".xlsx"), ("json", ".json")],
    )
    def test_the_filename_reflects_the_format(self, export_format: str, extension: str) -> None:
        exported = export_conversion(_conversion(), export_format)  # type: ignore[arg-type]

        assert exported.filename.endswith(extension)

    def test_an_unsafe_title_cannot_break_the_header(self) -> None:
        conversion = _conversion(title='weird<>:"name\nreport')

        filename = export_conversion(conversion, "md").filename

        for unsafe in ("<", ">", ":", '"', "\n", "\r", "/", "\\"):
            assert unsafe not in filename

    def test_an_empty_title_still_produces_a_filename(self) -> None:
        exported = export_conversion(_conversion(title=""), "md")

        assert exported.filename
        assert exported.filename.endswith(".md")


class TestDeterminism:
    @pytest.mark.parametrize("export_format", ["md", "json", "xlsx", "docx"])
    def test_repeated_export_is_stable(self, export_format: str) -> None:
        """Same stored conversion, same file — a download must not drift between clicks.

        DOCX and XLSX are ZIP containers whose entries carry timestamps, so the *bytes*
        can differ; what must not differ is the content, which is what is compared for
        those two.
        """
        conversion = _conversion()

        first = export_conversion(conversion, export_format)  # type: ignore[arg-type]
        second = export_conversion(conversion, export_format)  # type: ignore[arg-type]

        if export_format in {"md", "json"}:
            assert first.content == second.content
        elif export_format == "xlsx":
            assert list(
                load_workbook(io.BytesIO(first.content))["Customers"].iter_rows(values_only=True)
            ) == list(
                load_workbook(io.BytesIO(second.content))["Customers"].iter_rows(values_only=True)
            )
        else:
            assert [p.text for p in DocxDocument(io.BytesIO(first.content)).paragraphs] == [
                p.text for p in DocxDocument(io.BytesIO(second.content)).paragraphs
            ]

        assert first.filename == second.filename
        assert first.media_type == second.media_type
