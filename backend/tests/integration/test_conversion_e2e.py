"""End-to-end and hardening tests for document conversion (Phase 9).

**No test here calls a real API.** The project has a `live` marker for tests that hit real
services (deselected by `pytest.ini`'s `-m 'not live'`), and nothing in this file uses it:
everything runs through `FakeLLMProvider`, with real ingestion, a real database, real
repositories, and real export code.

Two jobs. First, one complete realistic run — a messy banking document uploaded, converted,
persisted, retrieved, and downloaded in all four formats — proving the whole chain works
together rather than only in parts. Second, hardening: what happens when the AI returns
something hostile or nonsensical, when values are missing, when text is not Latin, and
when someone asks for a result that is not theirs.

Earlier phases test their own units. What is exercised here is the seams between them, and
the failure modes that only appear once everything is connected.
"""

from __future__ import annotations

import io
import json
from dataclasses import dataclass, field
from types import SimpleNamespace
from typing import Any
from uuid import uuid4

import pytest
from docx import Document as DocxDocument
from fastapi import FastAPI
from httpx import ASGITransport, AsyncClient
from openpyxl import load_workbook
from pydantic import BaseModel

from app.ai.conversion_schema import (
    PlannedDocument,
    ProposedCorrectionList,
    ProposedCorrectionModel,
)
from app.ai.dynamic_schema import InferredStructureField, InferredStructureSchema
from app.api.deps import get_current_owner_id, get_llm_provider
from app.models.domain.insight import ContextBlock, ModelResult
from app.services.job_service import JobService
from tests.fixtures.builders import build_txt
from tests.fixtures.fake_provider import FakeLLMProvider

# A deliberately messy statement: three spellings of savings, a misspelled name, a
# customer with no balance recorded, and a non-Latin name.
BANKING_DOCUMENT = (
    "CUSTOMER MASTER — BRANCH 004\n\n"
    "Account holder Rajesh Kumr Sharma banks with State Bank of India. "
    "Account type: sav. KYC status: Pending. Balance: 145320.\n\n"
    "Account holder Priya Sharma banks with HDFC Bank. "
    "Account type: SAVINGS. KYC status: Completed. Balance: 45320.\n\n"
    "Account holder Rahul Verma banks with ICICI Bank. "
    "Account type: Current. KYC status: Pending. Balance: 3200.\n\n"
    "Account holder अनिता देसाई banks with Axis Bank. "
    "Account type: Savings. KYC status: Completed."
)

INSTRUCTION = (
    "Create a clean customer master report. Group customers by account type, "
    "standardize account types, put pending KYC issues in a separate section, "
    "preserve important details, and do not invent missing information."
)


@dataclass
class _SequencedFakeProvider(FakeLLMProvider):
    """Returns a different `.parsed` on each successive structured call, in order."""

    parsed_sequence: list[Any] = field(default_factory=list)

    async def generate_structured(
        self,
        *,
        system: str,
        instruction: str,
        context_blocks: list[ContextBlock],
        document_title: str,
        output_model: type[BaseModel],
    ) -> ModelResult:
        if self.parsed_sequence:
            self.parsed = self.parsed_sequence.pop(0)
        return await super().generate_structured(
            system=system,
            instruction=instruction,
            context_blocks=context_blocks,
            document_title=document_title,
            output_model=output_model,
        )


def _v(value: Any, quote: str) -> SimpleNamespace:
    return SimpleNamespace(value=value, source_quote=quote)


def _schema() -> InferredStructureSchema:
    return InferredStructureSchema(
        fields=[
            InferredStructureField(name="account_holder", type="string"),
            InferredStructureField(name="account_type", type="string"),
            InferredStructureField(name="kyc_status", type="string"),
            InferredStructureField(name="balance", type="number"),
        ]
    )


def _extraction() -> SimpleNamespace:
    return SimpleNamespace(
        records=[
            SimpleNamespace(
                account_holder=_v("Rajesh Kumr Sharma", "Account holder Rajesh Kumr Sharma"),
                account_type=_v("sav", "Account type: sav"),
                kyc_status=_v("Pending", "KYC status: Pending"),
                balance=_v(145320, "Balance: 145320"),
            ),
            SimpleNamespace(
                account_holder=_v("Priya Sharma", "Account holder Priya Sharma"),
                account_type=_v("SAVINGS", "Account type: SAVINGS"),
                kyc_status=_v("Completed", "KYC status: Completed"),
                balance=_v(45320, "Balance: 45320"),
            ),
            SimpleNamespace(
                account_holder=_v("Rahul Verma", "Account holder Rahul Verma"),
                account_type=_v("Current", "Account type: Current"),
                kyc_status=_v("Pending", "KYC status: Pending"),
                balance=_v(3200, "Balance: 3200"),
            ),
            # No balance stated for this customer — the absence must survive.
            SimpleNamespace(
                account_holder=_v("अनिता देसाई", "Account holder अनिता देसाई"),
                account_type=_v("Savings", "Account type: Savings"),
                kyc_status=_v("Completed", "KYC status: Completed"),
            ),
        ]
    )


def _plan(title: str = "Customer Master Report") -> PlannedDocument:
    return PlannedDocument.model_validate(
        {
            "title": title,
            "sections": [
                {
                    "heading": "Customers",
                    "kind": "table",
                    "columns": ["account_holder", "account_type", "balance"],
                    "group_by": "account_type",
                    "order_by": [{"field": "balance", "descending": True, "numeric": True}],
                },
                {
                    "heading": "Pending KYC",
                    "kind": "table",
                    "columns": ["account_holder", "kyc_status"],
                    "filters": [{"field": "kyc_status", "mode": "equals", "values": ["Pending"]}],
                },
                {"heading": "Notes", "kind": "note", "note": "Prepared for quarterly review."},
            ],
            "labels": [{"field": "account_holder", "label": "Account Holder"}],
            "vocabularies": [{"field": "account_type", "values": ["Savings", "Current", "Salary"]}],
        }
    )


def _proposals() -> ProposedCorrectionList:
    """Two provable fixes, one uncertain typo, and one confident hallucination."""
    return ProposedCorrectionList(
        corrections=[
            ProposedCorrectionModel(
                record_index=0,
                field="account_type",
                original="sav",
                corrected="Savings",
                kind="vocabulary",
                confidence="high",
                reason="standardized account type",
            ),
            ProposedCorrectionModel(
                record_index=1,
                field="account_type",
                original="SAVINGS",
                corrected="Savings",
                kind="vocabulary",
                confidence="high",
                reason="standardized account type",
            ),
            ProposedCorrectionModel(
                record_index=0,
                field="account_holder",
                original="Rajesh Kumr Sharma",
                corrected="Rajesh Kumar Sharma",
                kind="spelling",
                confidence="medium",
                reason="possible spelling correction",
            ),
            ProposedCorrectionModel(
                record_index=2,
                field="account_holder",
                original="Rahul Verma",
                corrected="Neha Kapoor",
                kind="spelling",
                confidence="high",
                reason="hallucinated rename",
            ),
        ]
    )


def _sequence(**kwargs: Any) -> list[Any]:
    return [
        kwargs.get("schema", _schema()),
        kwargs.get("extraction", _extraction()),
        kwargs.get("plan", _plan()),
        kwargs.get("proposals", _proposals()),
    ]


@pytest.fixture
def provider() -> _SequencedFakeProvider:
    return _SequencedFakeProvider(parsed_sequence=_sequence())


@pytest.fixture
def e2e_app(ai_app: FastAPI, provider: _SequencedFakeProvider) -> FastAPI:
    ai_app.dependency_overrides[get_llm_provider] = lambda: provider
    return ai_app


@pytest.fixture
async def client(e2e_app: FastAPI):  # type: ignore[no-untyped-def]
    transport = ASGITransport(app=e2e_app)
    async with (
        e2e_app.router.lifespan_context(e2e_app),
        AsyncClient(transport=transport, base_url="http://testserver") as http_client,
    ):
        yield http_client


@pytest.fixture
def jobs(e2e_app: FastAPI) -> JobService:
    return e2e_app.state.job_service  # type: ignore[no-any-return]


async def _ingest(
    client: AsyncClient, api_prefix: str, jobs: JobService, *, filename: str = "customers.txt"
) -> str:
    response = await client.post(
        f"{api_prefix}/documents",
        files={"file": (filename, build_txt(BANKING_DOCUMENT), "text/plain")},
    )
    assert response.status_code == 202, response.text
    document_id: str = response.json()["id"]
    await jobs.wait_for_all(timeout=30.0)
    assert (await client.get(f"{api_prefix}/documents/{document_id}")).json()["status"] == "ready"
    return document_id


async def _convert(client: AsyncClient, api_prefix: str, document_id: str, **body: Any) -> Any:
    return await client.post(
        f"{api_prefix}/documents/{document_id}/convert",
        json={"instruction": INSTRUCTION, **body},
    )


class TestCompleteWorkflow:
    """One realistic banking conversion, from upload through every export."""

    async def test_the_whole_chain(
        self,
        client: AsyncClient,
        api_prefix: str,
        jobs: JobService,
        provider: _SequencedFakeProvider,
    ) -> None:
        document_id = await _ingest(client, api_prefix, jobs)

        # --- convert -------------------------------------------------------
        created = (await _convert(client, api_prefix, document_id)).json()
        assert created["title"] == "Customer Master Report"
        assert created["source_record_count"] == 4

        # Organised as planned: grouped, filtered, and a note section.
        headings = [section["heading"] for section in created["sections"]]
        assert "Customers — Savings" in headings
        assert "Customers — Current" in headings
        assert "Pending KYC" in headings
        assert "Notes" in headings

        # Corrections: two applied, one suggestion, one refused.
        assert len(created["corrections"]) == 2
        assert len(created["suggestions"]) == 1
        assert len(created["rejected"]) == 1
        assert created["rejected"][0]["corrected"] == "Neha Kapoor"

        # Nothing invented anywhere in the rendered document.
        assert "Neha Kapoor" not in created["markdown"]
        assert "Rahul Verma" in created["markdown"]

        # --- retrieve ------------------------------------------------------
        fetched = (
            await client.get(f"{api_prefix}/documents/{document_id}/conversions/{created['id']}")
        ).json()
        assert fetched == created

        # --- export, every format -----------------------------------------
        calls_before_export = provider.call_count
        files: dict[str, bytes] = {}
        for export_format in ("md", "docx", "xlsx", "json"):
            response = await client.get(
                f"{api_prefix}/documents/{document_id}/conversions/{created['id']}"
                f"/export?format={export_format}"
            )
            assert response.status_code == 200, response.text
            files[export_format] = response.content

        # Not one AI call for any download.
        assert provider.call_count == calls_before_export

        assert files["md"].decode("utf-8") == created["markdown"]

        payload = json.loads(files["json"])
        assert payload["rejected"][0]["corrected"] == "Neha Kapoor"

        document = DocxDocument(io.BytesIO(files["docx"]))
        docx_text = "\n".join(p.text for p in document.paragraphs)
        assert "Changes & Corrections" in docx_text
        assert "Neha Kapoor" not in "\n".join(
            cell.text for table in document.tables for row in table.rows for cell in row.cells
        )

        workbook = load_workbook(io.BytesIO(files["xlsx"]))
        assert "Changes & Corrections" in workbook.sheetnames
        holders = {
            row[0]
            for name in workbook.sheetnames
            if name != "Changes & Corrections"
            for row in workbook[name].iter_rows(min_row=2, values_only=True)
        }
        assert "Rahul Verma" in holders
        assert "Neha Kapoor" not in holders

    async def test_the_customer_with_no_balance_keeps_that_absence(
        self, client: AsyncClient, api_prefix: str, jobs: JobService
    ) -> None:
        """The document never states a balance for one customer. No format may fill it."""
        document_id = await _ingest(client, api_prefix, jobs)
        created = (await _convert(client, api_prefix, document_id)).json()

        absent = [
            item
            for section in created["sections"]
            for record in section["records"]
            for item in record["fields"]
            if item["name"] == "balance" and item["value"]["origin"] == "absent"
        ]
        assert absent, "the record with no balance should carry an absent value"
        assert all(item["value"]["text"] == "" for item in absent)

        response = await client.get(
            f"{api_prefix}/documents/{document_id}/conversions/{created['id']}/export?format=xlsx"
        )
        workbook = load_workbook(io.BytesIO(response.content))
        for name in workbook.sheetnames:
            for row in workbook[name].iter_rows(min_row=2, values_only=True):
                assert all(cell in (None, "") or isinstance(cell, str | bool) for cell in row)


class TestUnicode:
    async def test_non_latin_values_survive_every_format(
        self, client: AsyncClient, api_prefix: str, jobs: JobService
    ) -> None:
        document_id = await _ingest(client, api_prefix, jobs)
        created = (await _convert(client, api_prefix, document_id)).json()

        assert "अनिता देसाई" in created["markdown"]

        for export_format in ("md", "json"):
            response = await client.get(
                f"{api_prefix}/documents/{document_id}/conversions/{created['id']}"
                f"/export?format={export_format}"
            )
            assert "अनिता देसाई" in response.content.decode("utf-8")

        docx = await client.get(
            f"{api_prefix}/documents/{document_id}/conversions/{created['id']}/export?format=docx"
        )
        cells = {
            cell.text
            for table in DocxDocument(io.BytesIO(docx.content)).tables
            for row in table.rows
            for cell in row.cells
        }
        assert "अनिता देसाई" in cells

        xlsx = await client.get(
            f"{api_prefix}/documents/{document_id}/conversions/{created['id']}/export?format=xlsx"
        )
        workbook = load_workbook(io.BytesIO(xlsx.content))
        values = {
            row[0]
            for name in workbook.sheetnames
            for row in workbook[name].iter_rows(min_row=2, values_only=True)
        }
        assert "अनिता देसाई" in values

    @pytest.mark.parametrize(
        "title", ["ग्राहक रिपोर्ट", "客户报告", "Rapport financiér", "Отчёт клиента"]
    )
    async def test_a_non_latin_title_does_not_break_the_download(
        self,
        client: AsyncClient,
        api_prefix: str,
        jobs: JobService,
        provider: _SequencedFakeProvider,
        title: str,
    ) -> None:
        """HTTP headers are latin-1. A report titled in Devanagari or Chinese must still
        download — RFC 6266 gives an ASCII fallback plus the real name, UTF-8 encoded."""
        provider.parsed_sequence = _sequence(plan=_plan(title=title))
        document_id = await _ingest(client, api_prefix, jobs)
        created = (await _convert(client, api_prefix, document_id)).json()

        response = await client.get(
            f"{api_prefix}/documents/{document_id}/conversions/{created['id']}/export?format=md"
        )

        assert response.status_code == 200, response.text
        disposition = response.headers["content-disposition"]
        disposition.encode("latin-1")  # must not raise
        assert "filename*=UTF-8''" in disposition
        assert "attachment" in disposition

    async def test_a_non_ascii_upload_filename_does_not_break_structure_export(
        self,
        client: AsyncClient,
        api_prefix: str,
        jobs: JobService,
        provider: _SequencedFakeProvider,
    ) -> None:
        """The same header defect reached the pre-existing structure export, whose
        filename comes straight from what the user uploaded."""
        provider.parsed_sequence = [
            _schema(),
            _extraction(),
        ]
        document_id = await _ingest(client, api_prefix, jobs, filename="客户报告.txt")

        structure = await client.post(
            f"{api_prefix}/documents/{document_id}/structure",
            json={"instruction": "One row per customer with holder and type."},
        )
        assert structure.status_code == 200, structure.text
        schema = structure.json()["schema_used"]

        response = await client.post(
            f"{api_prefix}/documents/{document_id}/structure/export?format=csv",
            json={"schema": schema},
        )

        assert response.status_code == 200, response.text
        response.headers["content-disposition"].encode("latin-1")  # must not raise


class TestHostileAiOutput:
    """The AI is untrusted. None of these may produce a wrong document."""

    async def test_a_plan_naming_fields_that_do_not_exist_is_stripped(
        self,
        client: AsyncClient,
        api_prefix: str,
        jobs: JobService,
        provider: _SequencedFakeProvider,
    ) -> None:
        hostile = PlannedDocument.model_validate(
            {
                "title": "Report",
                "sections": [
                    {
                        "heading": "Customers",
                        "kind": "table",
                        "columns": ["account_holder", "credit_score", "annual_income"],
                        "group_by": "branch_manager",
                    }
                ],
            }
        )
        provider.parsed_sequence = _sequence(plan=hostile)
        document_id = await _ingest(client, api_prefix, jobs)

        created = (await _convert(client, api_prefix, document_id)).json()

        assert created["sections"][0]["columns"] == ["account_holder"]
        body = created["markdown"].split("## Warnings")[0]
        for invented in ("credit_score", "annual_income", "branch_manager"):
            assert invented not in body

    async def test_corrections_that_change_numbers_are_refused(
        self,
        client: AsyncClient,
        api_prefix: str,
        jobs: JobService,
        provider: _SequencedFakeProvider,
    ) -> None:
        """A balance must never be edited, whatever the model claims."""
        hostile = ProposedCorrectionList(
            corrections=[
                ProposedCorrectionModel(
                    record_index=0,
                    field="balance",
                    original="145320",
                    corrected="999999",
                    kind="whitespace",
                    confidence="high",
                    reason="tidy up",
                ),
                ProposedCorrectionModel(
                    record_index=1,
                    field="kyc_status",
                    original="Completed",
                    corrected="Pending",
                    kind="case_normalization",
                    confidence="high",
                    reason="normalize",
                ),
            ]
        )
        provider.parsed_sequence = _sequence(proposals=hostile)
        document_id = await _ingest(client, api_prefix, jobs)

        created = (await _convert(client, api_prefix, document_id)).json()

        assert created["corrections"] == []
        assert len(created["rejected"]) == 2
        assert "999999" not in created["markdown"]
        assert "145320" in created["markdown"]

    async def test_a_correction_anchored_to_a_value_that_does_not_exist_is_refused(
        self,
        client: AsyncClient,
        api_prefix: str,
        jobs: JobService,
        provider: _SequencedFakeProvider,
    ) -> None:
        hostile = ProposedCorrectionList(
            corrections=[
                ProposedCorrectionModel(
                    record_index=0,
                    field="account_holder",
                    original="Someone Who Is Not There",
                    corrected="Anyone At All",
                    kind="spelling",
                    confidence="high",
                    reason="anchored to nothing",
                )
            ]
        )
        provider.parsed_sequence = _sequence(proposals=hostile)
        document_id = await _ingest(client, api_prefix, jobs)

        created = (await _convert(client, api_prefix, document_id)).json()

        assert created["corrections"] == []
        assert "anchor mismatch" in created["rejected"][0]["rejection_reason"]
        assert "Anyone At All" not in created["markdown"]

    async def test_a_formula_in_a_source_value_cannot_execute_in_the_spreadsheet(
        self,
        client: AsyncClient,
        api_prefix: str,
        jobs: JobService,
        provider: _SequencedFakeProvider,
    ) -> None:
        """The value came from the user's document, so it is untrusted like any other
        document content."""
        hostile_extraction = SimpleNamespace(
            records=[
                SimpleNamespace(
                    account_holder=_v("=cmd|' /C calc'!A0", "Account holder Rajesh Kumr Sharma"),
                    account_type=_v("Savings", "Account type: Savings"),
                    kyc_status=_v("Completed", "KYC status: Completed"),
                    balance=_v(100, "Balance: 145320"),
                )
            ]
        )
        provider.parsed_sequence = _sequence(extraction=hostile_extraction)
        document_id = await _ingest(client, api_prefix, jobs)
        created = (await _convert(client, api_prefix, document_id)).json()

        response = await client.get(
            f"{api_prefix}/documents/{document_id}/conversions/{created['id']}/export?format=xlsx"
        )

        workbook = load_workbook(io.BytesIO(response.content))
        cells = [
            cell
            for name in workbook.sheetnames
            for row in workbook[name].iter_rows(values_only=True)
            for cell in row
            if isinstance(cell, str)
        ]
        assert any(cell.startswith("'=cmd") for cell in cells)
        assert not any(cell.startswith("=cmd") for cell in cells)

    async def test_an_empty_plan_is_reported_not_faked(
        self,
        client: AsyncClient,
        api_prefix: str,
        jobs: JobService,
        provider: _SequencedFakeProvider,
    ) -> None:
        provider.parsed_sequence = _sequence(
            plan=PlannedDocument.model_validate({"title": "Report", "sections": []})
        )
        document_id = await _ingest(client, api_prefix, jobs)

        response = await _convert(client, api_prefix, document_id)

        assert response.status_code == 500, response.text
        assert response.json()["error"]["code"] == "invalid_model_response"

    async def test_an_extraction_that_finds_nothing_is_reported(
        self,
        client: AsyncClient,
        api_prefix: str,
        jobs: JobService,
        provider: _SequencedFakeProvider,
    ) -> None:
        provider.parsed_sequence = [_schema(), SimpleNamespace(records=[])]
        document_id = await _ingest(client, api_prefix, jobs)

        response = await _convert(client, api_prefix, document_id)

        assert response.status_code == 422, response.text
        assert response.json()["error"]["code"] == "insufficient_content"


class TestRegressionOfExistingFeatures:
    """Conversion must not have disturbed what already worked."""

    async def test_summary_still_works(
        self,
        client: AsyncClient,
        api_prefix: str,
        jobs: JobService,
        provider: _SequencedFakeProvider,
    ) -> None:
        document_id = await _ingest(client, api_prefix, jobs)
        provider.parsed_sequence = []
        provider.text = "A short summary of the customer master."

        response = await client.post(f"{api_prefix}/documents/{document_id}/summary")

        assert response.status_code == 200, response.text
        assert response.json()["kind"] == "summary"

    async def test_structure_still_returns_records_and_schema(
        self,
        client: AsyncClient,
        api_prefix: str,
        jobs: JobService,
        provider: _SequencedFakeProvider,
    ) -> None:
        provider.parsed_sequence = [_schema(), _extraction()]
        document_id = await _ingest(client, api_prefix, jobs)

        response = await client.post(
            f"{api_prefix}/documents/{document_id}/structure",
            json={"instruction": "One row per customer."},
        )

        assert response.status_code == 200, response.text
        body = response.json()
        assert body["record_count"] == 4
        assert "account_holder" in body["schema_used"]

    async def test_a_conversion_does_not_appear_in_the_insights_listing(
        self, client: AsyncClient, api_prefix: str, jobs: JobService
    ) -> None:
        """Conversions live in their own table; `/insights` must not start reporting a
        kind the existing clients do not know."""
        document_id = await _ingest(client, api_prefix, jobs)
        await _convert(client, api_prefix, document_id)

        insights = (await client.get(f"{api_prefix}/documents/{document_id}/insights")).json()

        assert all(
            item["kind"] in {"summary", "explanation", "extraction"} for item in insights["items"]
        )


class TestOwnershipUnderPressure:
    async def test_no_conversion_endpoint_leaks_across_owners(
        self, e2e_app: FastAPI, client: AsyncClient, api_prefix: str, jobs: JobService
    ) -> None:
        """One sweep over every conversion route: knowing an id must never be enough."""
        document_id = await _ingest(client, api_prefix, jobs)
        created = (await _convert(client, api_prefix, document_id)).json()

        e2e_app.dependency_overrides[get_current_owner_id] = lambda: "intruder"
        try:
            get_one = await client.get(
                f"{api_prefix}/documents/{document_id}/conversions/{created['id']}"
            )
            listing = await client.get(f"{api_prefix}/documents/{document_id}/conversions")
            export = await client.get(
                f"{api_prefix}/documents/{document_id}/conversions/{created['id']}"
                "/export?format=json"
            )
            convert = await _convert(client, api_prefix, document_id)
        finally:
            e2e_app.dependency_overrides.pop(get_current_owner_id, None)

        assert get_one.status_code == 404
        assert export.status_code == 404
        assert convert.status_code == 404
        assert listing.json()["items"] == []

    async def test_a_conversion_id_from_another_document_is_not_found(
        self, client: AsyncClient, api_prefix: str, jobs: JobService
    ) -> None:
        document_id = await _ingest(client, api_prefix, jobs)
        created = (await _convert(client, api_prefix, document_id)).json()

        response = await client.get(
            f"{api_prefix}/documents/{uuid4()}/conversions/{created['id']}/export?format=md"
        )

        assert response.status_code == 404


class TestDeterminismEndToEnd:
    async def test_the_same_decisions_produce_the_same_document(
        self, e2e_app: FastAPI, client: AsyncClient, api_prefix: str, jobs: JobService
    ) -> None:
        """With the model's replies fixed, the whole pipeline is reproducible."""
        document_id = await _ingest(client, api_prefix, jobs)
        first = (await _convert(client, api_prefix, document_id, refresh=True)).json()

        fresh = _SequencedFakeProvider(parsed_sequence=_sequence())
        e2e_app.dependency_overrides[get_llm_provider] = lambda: fresh
        second = (await _convert(client, api_prefix, document_id, refresh=True)).json()

        assert first["markdown"] == second["markdown"]
        assert first["sections"] == second["sections"]
        assert first["id"] != second["id"]  # still two distinct runs
